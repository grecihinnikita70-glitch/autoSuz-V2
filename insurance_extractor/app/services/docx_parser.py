from pathlib import Path

from docx import Document as DocxDocument


def extract_text_from_docx(file_path: str | Path) -> str:
    """Read paragraphs and tables from a DOCX file as plain text."""
    docx_document = DocxDocument(str(file_path))
    lines: list[str] = []

    for paragraph in docx_document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    # Insurance contracts often contain key data in tables. For the MVP we only
    # convert each row to readable text; structured field extraction comes later.
    for table in docx_document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)
